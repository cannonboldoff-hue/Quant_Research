"""Build the journal manuscript (.docx, .html, .pdf) from manuscript.md.

Tables are not typed into the manuscript: every ``{{tbl:NAME}}`` placeholder is
replaced by a table generated here from the CSV outputs of ``scripts/tfml_report.py``
(``paper/tables``), so the manuscript stays consistent with the experiment registry.

    python paper/manuscript/build_manuscript.py            # docx + html + pdf
    python paper/manuscript/build_manuscript.py --no-pdf   # skip headless-Chrome PDF

Needs ``pypandoc_binary`` and ``python-docx`` (see requirements-tfml.txt); the PDF is
printed by a local Chrome or Edge.
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
TABLES = PAPER / "tables"
FIGURES = PAPER / "figures"
OUT_STEM = "ML_Trend_Following_manuscript"

STAGES = ["S1_filter", "S2_regime", "S3_sizing", "S4_exit", "S5_params"]
STAGE_LABEL = {
    "S1_filter": "S1 Filter",
    "S2_regime": "S2 Regime",
    "S3_sizing": "S3 Sizing",
    "S4_exit": "S4 Exit",
    "S5_params": "S5 Params",
}
MODEL_LABEL = {
    "logreg": "Logistic regression",
    "rf": "Random forest",
    "et": "Extra trees",
    "lgbm": "LightGBM",
    "mlp": "MLP",
}
MARKET_LABEL = {
    "equity_index": "Equity indices",
    "etf": "ETFs",
    "futures": "Futures",
    "fx_major": "FX majors (G10)",
    "fx_cross": "FX crosses (G10)",
    "fx_em": "FX emerging",
    "crypto": "Crypto",
}
FAMILY_LABEL = {
    "ma_crossover": "MA crossover",
    "price_vs_ma": "Price vs MA",
    "adaptive_ma": "Adaptive MA",
    "breakout": "Channel breakout",
    "volatility_stop": "Volatility stop",
    "time_series_momentum": "Time-series momentum",
    "oscillator_trend": "Oscillator trend",
    "directional": "Directional movement",
}
RUN_LABEL = {
    "benchmark": "Benchmark (expanding WF, daily)",
    "robust_rolling10y": "Rolling 10-year window",
    "robust_holdout2013": "Single hold-out (2013)",
    "robust_single_stocks": "Single stocks (survivorship-biased)",
    "crypto_intraday_4h": "Crypto, 4-hour bars",
    "crypto_intraday_1h": "Crypto, 1-hour bars",
}
BLOCK_LABEL = {
    "all": "All 56 indicators",
    "momentum": "Momentum only",
    "volatility": "Volatility only",
    "trend": "Trend only",
    "oscillator_channel": "Oscillator/channel only",
    "other": "Other only",
    "context_only": "No indicators (context only)",
}
GROUP_LABEL = {
    "equity_index": "Equity indices",
    "futures": "Continuous futures",
    "fx_major": "FX majors",
    "fx_cross": "FX crosses",
    "fx_em": "FX emerging markets",
    "crypto": "Cryptocurrencies",
    "etf": "ETFs",
    "crypto_binance_intraday": "Crypto (Binance spot)",
    "us_stocks_sp100_current": "US stocks (current S&P 100)",
    "india_stocks_nifty_current": "Indian stocks (current NIFTY 50)",
}


def rd(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES / f"{name}.csv")


def f2(x, signed: bool = False) -> str:
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "–"
    s = f"{x:+.2f}" if signed else f"{x:.2f}"
    return s.replace("-", "−")


def f3(x) -> str:
    return "–" if not np.isfinite(x) else f"{x:.3f}".replace("-", "−")


def pct(x, nd: int = 1, signed: bool = False) -> str:
    if not np.isfinite(x):
        return "–"
    s = f"{100 * x:+.{nd}f}" if signed else f"{100 * x:.{nd}f}"
    return s.replace("-", "−")


def pval(p) -> str:
    if not np.isfinite(p):
        return "–"
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def md_table(header: list[str], rows: list[list[str]], align: str | None = None) -> str:
    """Pipe table; ``align`` is a string of l/r/c per column (default: first l, rest r)."""
    align = align or ("l" + "r" * (len(header) - 1))
    sep = {"l": ":---", "r": "---:", "c": ":---:"}
    out = ["| " + " | ".join(header) + " |", "|" + "|".join(sep[a] for a in align) + "|"]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


# ---------------------------------------------------------------- tables
def t_data() -> str:
    d = rd("T1_data_coverage")
    order = ["equity_index", "futures", "fx_major", "fx_cross", "fx_em", "crypto", "etf",
             "crypto_binance_intraday", "us_stocks_sp100_current", "india_stocks_nifty_current"]
    freq_order = {"daily": 0, "4h": 1, "1h": 2, "daily_binance": 3}
    d = d.assign(_g=d.group.map(order.index), _f=d.freq.map(freq_order)).sort_values(["_g", "_f"])
    rows = []
    for _, r in d.iterrows():
        role = "Benchmark" if str(r["benchmark"]) == "True" else "Robustness only"
        if r["survivorship"] == "current_constituents":
            role = "Robustness only (survivorship-biased)"
        rows.append([
            GROUP_LABEL.get(r["group"], r["group"]),
            {"daily_binance": "daily"}.get(r["freq"], r["freq"]),
            int(r["instruments"]),
            str(r["first_start"])[:4],
            f"{int(r['total_bars']):,}",
            r["source"].capitalize(),
            role,
        ])
    return md_table(["Group", "Bars", "Instr.", "From", "Bars (total)", "Source", "Use"], rows, "lcrrrll")


def t_headline() -> str:
    d = rd("benchmark_T5_headline_by_stage").set_index("stage")
    rows = []
    for s in STAGES:
        r = d.loc[s]
        rows.append([
            STAGE_LABEL[s], f2(r.mean_dSharpe, True), f2(r.median_dSharpe, True),
            pct(r.share_improved, 0), int(r.sig_pos_raw), int(r.sig_pos_holm), int(r.sig_pos_bh),
            int(r.sig_neg_bh), pct(r.mean_d_maxdd, 1, True), f2(r.mean_d_turnover, True),
            pval(r.sign_test_p),
        ])
    return md_table(
        ["Stage", "Mean ΔSR", "Median ΔSR", "Improved (%)", "Sig.+ raw", "Sig.+ Holm",
         "Sig.+ BH", "Sig.− BH", "ΔMaxDD (pp)", "ΔTurnover", "Sign test p"], rows)


def t_composite() -> str:
    d = rd("benchmark_T17_composite_portfolio").set_index("config")
    rows = []
    for c in ["base"] + STAGES:
        r = d.loc[c]
        ci = "–" if c == "base" else f"[{f2(r.ci_lo)}, {f2(r.ci_hi)}]"
        rows.append([
            "Rule-based" if c == "base" else STAGE_LABEL[c], pct(r.cagr, 2), pct(r.ann_vol, 2),
            f2(r.sharpe), f2(r.sortino), f2(r.calmar), pct(r.max_drawdown, 1),
            pct(r.pct_pos_years, 0), "–" if c == "base" else f2(r.d_sharpe, True), ci,
            "–" if c == "base" else pval(r.p_boot_two),
        ])
    return md_table(["Configuration", "CAGR (%)", "Vol. (%)", "Sharpe", "Sortino", "Calmar",
                     "MaxDD (%)", "Pos. years (%)", "ΔSR", "95% CI", "p"], rows)


def t_models() -> str:
    d = rd("benchmark_T8_model_comparison")
    sel = rd("benchmark_T8c_selected_vs_alternatives").set_index("ml_stage")
    rows = []
    for m in ["logreg", "rf", "et", "lgbm", "mlp"]:
        row = [MODEL_LABEL[m]]
        for s in STAGES:
            r = d[(d.ml_model == m) & (d.ml_stage == s)].iloc[0]
            row.append(f"{f2(r.mean_dSharpe, True)} ({pct(r.selection_share, 0)}%)")
        rows.append(row)
    for lab, col in [("**Validation-selected**", "selected_model"), ("Average of five models", "avg_model"),
                     ("Best model ex post", "best_ex_post"), ("Worst model ex post", "worst_ex_post")]:
        rows.append([lab] + [f2(sel.loc[s, col], True) for s in STAGES])
    return md_table(["Model"] + [STAGE_LABEL[s] for s in STAGES], rows)


def t_markets() -> str:
    d = rd("benchmark_T9_market_groups")
    rows = []
    for m in ["equity_index", "etf", "futures", "fx_em", "fx_cross", "fx_major", "crypto"]:
        sub = d[d.market == m].set_index("ml_stage")
        row = [MARKET_LABEL[m], f2(sub.mean_base_sharpe.iloc[0])]
        row += [f"{f2(sub.loc[s, 'mean_dSharpe'], True)} ({int(sub.loc[s, 'sig_pos_bh'])})" for s in STAGES]
        rows.append(row)
    return md_table(["Market group", "Base SR"] + [STAGE_LABEL[s] for s in STAGES], rows)


def t_families() -> str:
    d = rd("benchmark_T13_strategy_family")
    order = d[d.ml_stage == "S1_filter"].sort_values("mean_base_sharpe").strategy_family
    rows = []
    for fam in order:
        sub = d[d.strategy_family == fam].set_index("ml_stage")
        row = [FAMILY_LABEL[fam], int(sub.n.iloc[0]), f2(sub.mean_base_sharpe.iloc[0])]
        row += [f2(sub.loc[s, "mean_dSharpe"], True) for s in STAGES]
        rows.append(row)
    return md_table(["Strategy family", "n", "Base SR"] + [STAGE_LABEL[s] for s in STAGES], rows)


TERM_LABEL = {
    "Intercept": "Intercept (S1, logistic regression, MA crossover, equity)",
    "base_sharpe": "Baseline Sharpe ratio of the instrument",
}


def _term(t: str) -> str:
    if t in TERM_LABEL:
        return TERM_LABEL[t]
    k, v = t.split("=", 1)
    lab = {**STAGE_LABEL, **MODEL_LABEL, **FAMILY_LABEL, "crypto": "Crypto", "etf": "ETF",
           "futures": "Futures", "fx": "FX"}.get(v, v)
    kind = {"ml_stage": "Stage", "ml_model": "Model", "strategy_family": "Family",
            "asset_class": "Asset class"}[k]
    return f"{kind}: {lab}"


def t_regression() -> str:
    d = rd("benchmark_T15_conditions_regression")
    rows = [[_term(r.term), f3(r.coef), f3(r.se), f2(r.t), pval(r.p)] for r in d.itertuples()]
    return md_table(["Term", "Coef.", "SE", "t", "p"], rows)


def t_subperiods() -> str:
    d = rd("benchmark_T11_subperiods")
    rows = []
    for p in sorted(d.period.unique()):
        sub = d[d.period == p].set_index("ml_stage")
        row = [p.replace("-", "–"), f2(sub.mean_base_sharpe.iloc[0])]
        row += [f"{f2(sub.loc[s, 'mean_dSharpe'], True)} ({int(sub.loc[s, 'sig_pos'])})" for s in STAGES]
        rows.append(row)
    return md_table(["Period", "Base SR"] + [STAGE_LABEL[s] for s in STAGES], rows)


BT_LABEL = {
    "next_open|cost1|vol_target0.1 (benchmark)": "Benchmark: next open, 1× costs, 10% vol target",
    "next_open|cost0|vol_target0.1": "Zero costs",
    "next_open|cost3|vol_target0.1": "3× costs",
    "next_close|cost1|vol_target0.1": "Next-close execution (one bar later)",
    "same_close|cost1|vol_target0.1": "Same-close execution (optimistic)",
    "next_open|cost1|unit": "Unit sizing (no vol targeting)",
}


def t_sensitivity() -> str:
    d = rd("benchmark_T12_backtest_sensitivity")
    rows = []
    for k, lab in BT_LABEL.items():
        sub = d[d.backtest == k].set_index("ml_stage")
        row = [lab, f2(sub.mean_base_sharpe.iloc[0])]
        row += [f"{f2(sub.loc[s, 'mean_dSharpe'], True)} ({int(sub.loc[s, 'sig_pos'])}/{int(sub.loc[s, 'sig_neg'])})"
                for s in STAGES]
        rows.append(row)
    return md_table(["Backtest assumption", "Base SR"] + [STAGE_LABEL[s] for s in STAGES], rows)


def t_placebo() -> str:
    d = rd("benchmark_T19_placebo_by_stage").set_index("stage")
    rows = []
    for s in STAGES:
        r = d.loc[s]
        rows.append([STAGE_LABEL[s], f2(r.mean_d_ml, True), f2(r.mean_d_placebo, True),
                     f2(r.mean_ml_minus_placebo, True), pct(r.share_ml_beats_placebo_mean, 0),
                     int(r.n_ml_above_placebo_p95), int(r.n_bh)])
    return md_table(["Stage", "ΔSR, ML", "ΔSR, placebo", "ML − placebo", "ML > placebo mean (%)",
                     "ML > placebo 95th pct.", "Same, BH"], rows)


def t_null() -> str:
    d = rd("T21_null_randomwalk").set_index("stage")
    rows = []
    for s in STAGES:
        r = d.loc[s]
        rows.append([STAGE_LABEL[s], f2(r.mean_dSharpe_selected, True), f2(r.mean_dSharpe_all_models, True),
                     pct(r.share_improved, 0), int(r.sig_pos_raw), int(r.sig_neg_raw),
                     pct(r.instrument_false_pos_rate, 1), f2(r.mean_dSharpe_cost0, True),
                     f2(r.mean_dSharpe_cost3, True)])
    return md_table(["Stage", "ΔSR selected", "ΔSR all models", "Improved (%)", "Sig.+", "Sig.−",
                     "Instr. false-pos. (%)", "ΔSR 0× cost", "ΔSR 3× cost"], rows)


def t_ablation() -> str:
    d = rd("T20_feature_family_ablation")
    d = d[d.model == "selected"]
    rows = []
    for b, lab in BLOCK_LABEL.items():
        sub = d[d.feature_block == b].set_index("stage")
        row = [lab]
        for s in ["S1_filter", "S3_sizing", "S4_exit"]:
            r = sub.loc[s]
            row.append(f"{f2(r.mean_dSharpe, True)} ({int(r.sig_pos)}/{int(r.n)})")
        rows.append(row)
    return md_table(["Feature set seen by the model", "S1 Filter", "S3 Sizing", "S4 Exit"], rows)


def t_runs() -> str:
    d = rd("T18_cross_run_robustness")
    rows = []
    for run, lab in RUN_LABEL.items():
        sub = d[d.run == run].set_index("stage")
        if sub.empty:
            continue
        row = [lab, f2(sub.mean_base_sharpe.iloc[0])]
        row += [f"{f2(sub.loc[s, 'mean_dSharpe'], True)} ({int(sub.loc[s, 'sig_pos_bh'])}/{int(sub.loc[s, 'sig_neg_bh'])})"
                for s in STAGES]
        rows.append(row)
    return md_table(["Run", "Base SR"] + [STAGE_LABEL[s] for s in STAGES], rows)


def t_strategies() -> str:
    d = rd("T2_strategies")
    rows = [[f"`{r.id}`", r.name, FAMILY_LABEL[r.family], r.reference, r.params.replace("_=", "="),
             r.direction] for r in d.itertuples()]
    return md_table(["ID", "Rule", "Family", "Reference", "Parameters", "Direction"], rows, "llllll")


def t_features() -> str:
    d = rd("T3_features")
    rows = [[f"`{r.feature}`", r.family.replace("_", " "), r.definition] for r in d.itertuples()]
    return md_table(["Feature", "Family", "Definition"], rows, "lll")


def t_baseline_metrics() -> str:
    d = rd("benchmark_T7_baseline_metrics").sort_values("sharpe", ascending=False)
    rows = [[f"`{r.strategy}`", pct(r.cagr, 2), pct(r.ann_vol, 1), f2(r.sharpe), f2(r.sortino),
             f2(r.calmar), pct(r.max_drawdown, 1), pct(r.win_rate_trades, 0), f2(r.profit_factor_trades),
             f"{r.turnover:.1f}", f"{int(r.n_trades):,}", pct(r.pct_pos_years, 0)] for r in d.itertuples()]
    return md_table(["Strategy", "CAGR (%)", "Vol. (%)", "Sharpe", "Sortino", "Calmar", "MaxDD (%)",
                     "Win rate (%)", "Profit factor", "Turnover", "Trades", "Pos. years (%)"], rows)


def t_strategy_ml() -> str:
    base = rd("benchmark_T7_baseline_metrics").set_index("strategy")
    det = rd("benchmark_T6_strategy_stage_detail")
    spa = rd("benchmark_T14_spa_reality_check").set_index("strategy")
    pbo = rd("benchmark_T16_pbo_dsr").set_index("strategy")
    rows = []
    for s in base.sort_values("sharpe", ascending=False).index:
        row = [f"`{s}`", f2(base.loc[s, "sharpe"])]
        for st in STAGES:
            r = det[(det.strategy == s) & (det.ml_stage == st)].iloc[0]
            star = "*" if r.p_bh < 0.05 else ""
            row.append(f"{f2(r.d_sharpe, True)}{star}")
        row += [pval(spa.loc[s, "p_spa"]), f3(pbo.loc[s, "pbo"]), f2(pbo.loc[s, "dsr_best"])]
        rows.append(row)
    return md_table(["Strategy", "Base SR"] + [STAGE_LABEL[s] for s in STAGES] + ["SPA p", "PBO", "DSR"], rows)


def t_validation() -> str:
    d = rd("T0_data_validation")
    rows = [[r.pair, f"{int(r.overlap_days):,}", f2(r.corr_daily), f2(r.corr_weekly),
             pct(r.tracking_err_ann, 1)] for r in d.itertuples()]
    return md_table(["Comparison", "Overlap (days)", "Corr. daily", "Corr. weekly", "Tracking error (%/yr)"], rows)


TABLES_FN = {
    "data": t_data, "headline": t_headline, "composite": t_composite, "models": t_models,
    "markets": t_markets, "families": t_families, "regression": t_regression,
    "subperiods": t_subperiods, "sensitivity": t_sensitivity, "placebo": t_placebo,
    "null": t_null, "ablation": t_ablation, "runs": t_runs, "strategies": t_strategies,
    "features": t_features, "baseline_metrics": t_baseline_metrics,
    "strategy_ml": t_strategy_ml, "validation": t_validation,
}


def make_ablation_figure(path: Path) -> None:
    """Figure 7 redrawn from T20 with readable labels (same numbers as paper/figures/F7)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    d = rd("T20_feature_family_ablation")
    d = d[d.model == "selected"]
    blocks = list(BLOCK_LABEL)
    labels = ["All 56\nindicators", "Momentum\nonly", "Volatility\nonly", "Trend\nonly",
              "Oscillator/\nchannel only", "Other\nonly", "No indicators\n(context only)"]
    stages = [("S1_filter", "S1 trade filter", "#2b6cb0"), ("S3_sizing", "S3 sizing", "#2f9e6e"),
              ("S4_exit", "S4 exit", "#d69e2e")]
    x = np.arange(len(blocks))
    fig, ax = plt.subplots(figsize=(8.0, 3.6), dpi=200)
    for k, (s, lab, col) in enumerate(stages):
        vals = [d[(d.feature_block == b) & (d.stage == s)].mean_dSharpe.iloc[0] for b in blocks]
        ax.bar(x + (k - 1) * 0.26, vals, 0.26, label=lab, color=col)
    ax.axvspan(-0.5, 0.5, color="#000", alpha=0.05, lw=0)
    ax.axvspan(len(blocks) - 1.5, len(blocks) - 0.5, color="#000", alpha=0.05, lw=0)
    ax.set_xticks(x, labels, fontsize=8)
    ax.set_xlim(-0.5, len(blocks) - 0.5)
    ax.set_ylabel("Mean ΔSharpe (selected model)", fontsize=8.5)
    ax.tick_params(axis="y", labelsize=8)
    ax.axhline(0, color="#000", lw=0.6)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(fontsize=8, frameon=False, ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.12))
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def render(src: str) -> str:
    def sub(m):
        return TABLES_FN[m.group(1)]()
    out = re.sub(r"\{\{tbl:([a-z_]+)\}\}", sub, src)
    left = re.findall(r"\{\{[^}]*\}\}", out)
    if left:
        raise SystemExit(f"unresolved placeholders: {left}")
    return out


# ---------------------------------------------------------------- docx styling
def make_reference_docx(pandoc: str, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    subprocess.run([pandoc, "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(str(path))
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin = sec.bottom_margin = Cm(2.3)
    st = doc.styles
    black = RGBColor(0, 0, 0)

    def font(name, size, bold=None, italic=None, color=black):
        s = st[name]
        s.font.name, s.font.size = "Cambria", Pt(size)
        if bold is not None:
            s.font.bold = bold
        if italic is not None:
            s.font.italic = italic
        s.font.color.rgb = color
        rpr = s.element.get_or_add_rPr()
        rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if rfonts is not None:
            for a in ("ascii", "hAnsi", "cs", "eastAsia"):
                rfonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{a}", "Cambria")
            for a in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
                rfonts.attrib.pop(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{a}", None)
        return s

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [x.name for x in st]:
            s = font(name, 11)
            s.paragraph_format.line_spacing = 1.15
            s.paragraph_format.space_after = Pt(6)
            if name != "Compact":
                s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    t = font("Title", 17, bold=True)
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Subtitle", "Author", "Date"):
        if name in [x.name for x in st]:
            s = font(name, 11)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Abstract" in [x.name for x in st]:
        a = font("Abstract", 10)
        a.paragraph_format.left_indent = a.paragraph_format.right_indent = Cm(1.0)
        a.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for lvl, size in ((1, 13), (2, 11.5), (3, 11)):
        h = font(f"Heading {lvl}", size, bold=True, italic=(lvl == 3))
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
        h.paragraph_format.space_after = Pt(4)
    for name in ("Table Caption", "Image Caption", "Caption"):
        if name in [x.name for x in st]:
            c = font(name, 9.5, italic=False)
            c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            c.paragraph_format.keep_with_next = True
    if "Block Text" in [x.name for x in st]:
        b = font("Block Text", 10.5, italic=True)
        b.paragraph_format.left_indent = b.paragraph_format.right_indent = Cm(0.8)
    if "Source Code" in [x.name for x in st]:
        st["Source Code"].font.size = Pt(9)
    doc.save(str(path))


def shrink_docx_tables(path: Path) -> None:
    """Small font, compact rows and a booktabs-like look for every table."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document(str(path))
    for tbl in doc.tables:
        ncols = len(tbl.columns)
        size = Pt(8.5 if ncols <= 7 else 7.5 if ncols <= 10 else 6.5)
        tblPr = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge, val in (("top", "single"), ("bottom", "single"), ("left", "nil"), ("right", "nil"),
                          ("insideH", "nil"), ("insideV", "nil")):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), val)
            if val == "single":
                el.set(qn("w:sz"), "8")
                el.set(qn("w:color"), "000000")
            borders.append(el)
        old = tblPr.find(qn("w:tblBorders"))
        if old is not None:
            tblPr.remove(old)
        tblPr.append(borders)
        for i, row in enumerate(tbl.rows):
            for cell in row.cells:
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcb = OxmlElement("w:tcBorders")
                    b = OxmlElement("w:bottom")
                    b.set(qn("w:val"), "single")
                    b.set(qn("w:sz"), "6")
                    b.set(qn("w:color"), "000000")
                    tcb.append(b)
                    tcPr.append(tcb)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.size = size
                        if i == 0:
                            r.font.bold = True
    doc.save(str(path))


# ---------------------------------------------------------------- pdf
CSS = r"""
@page { size: A4; margin: 22mm 20mm 22mm 20mm;
  @bottom-center { content: counter(page); font-family: Cambria, Georgia, serif; font-size: 9pt; } }
html { font-family: Cambria, "Times New Roman", Georgia, serif; font-size: 10.5pt; line-height: 1.42; color: #000; background: #fff; }
body { max-width: none; margin: 0; padding: 0; }
header#title-block-header { text-align: center; margin-bottom: 1.2em; }
h1.title { font-size: 17pt; line-height: 1.25; margin: 0 0 .5em 0; }
p.author, p.date { margin: .15em 0; font-size: 10.5pt; }
div.abstract { margin: 1.2em 2.2em 0 2.2em; font-size: 9.6pt; text-align: justify; }
div.abstract-title { font-weight: bold; text-align: center; margin-bottom: .3em; }
h1 { font-size: 13pt; margin: 1.4em 0 .4em 0; page-break-after: avoid; }
h2 { font-size: 11.2pt; margin: 1.1em 0 .3em 0; page-break-after: avoid; }
h3 { font-size: 10.5pt; font-style: italic; margin: .9em 0 .25em 0; page-break-after: avoid; }
p { text-align: justify; margin: .35em 0 .55em 0; hyphens: auto; orphans: 3; widows: 3; }
blockquote { margin: .6em 2em; font-style: italic; }
table { border-collapse: collapse; margin: .4em auto 1.0em auto; font-size: 8.2pt; line-height: 1.25;
  border-top: 1.2px solid #000; border-bottom: 1.2px solid #000; page-break-inside: auto; }
thead th { border-bottom: .8px solid #000; font-weight: bold; padding: 3px 5px; vertical-align: bottom; }
td { padding: 2px 5px; vertical-align: top; }
tr { page-break-inside: avoid; }
caption { caption-side: top; text-align: left; font-size: 9pt; margin-bottom: .35em; line-height: 1.3; }
figure { margin: .8em 0 1.1em 0; text-align: center; page-break-inside: avoid; }
figure img { max-width: 92%; max-height: 135mm; }
.keep { break-inside: avoid; page-break-inside: avoid; }
figcaption { font-size: 9pt; text-align: left; margin: .35em 4% 0 4%; line-height: 1.3; }
code { font-family: Consolas, "Courier New", monospace; font-size: 88%; }
pre { font-size: 8.5pt; background: #f6f6f6; padding: .5em .8em; page-break-inside: avoid; }
.wide table { font-size: 7.1pt; }
.small { font-size: 9pt; }
div.notes p { font-size: 8.6pt; margin-top: -.6em; }
#refs div.csl-entry, .references p { padding-left: 1.6em; text-indent: -1.6em; text-align: left; margin: .15em 0; font-size: 9.4pt; }
math { font-size: 105%; }
.pagebreak { page-break-before: always; }
"""


def find_browser() -> str | None:
    for p in (r"C:\Program Files\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if Path(p).exists():
            return p
    for name in ("google-chrome", "chromium", "chrome", "msedge"):
        if shutil.which(name):
            return shutil.which(name)
    return None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()
    import pypandoc

    pandoc = pypandoc.get_pandoc_path()
    src = (HERE / "manuscript.md").read_text(encoding="utf-8")
    body = render(src)
    built = HERE / "build"
    built.mkdir(exist_ok=True)
    make_ablation_figure(built / "F7_ablation.png")
    md = built / f"{OUT_STEM}.md"
    md.write_text(body, encoding="utf-8")
    common = ["--from", "markdown+pipe_tables+tex_math_dollars+fenced_divs+bracketed_spans+implicit_figures",
              "--resource-path", f"{HERE}{';' if sys.platform == 'win32' else ':'}{PAPER}",
              "--citeproc", "--bibliography", str(HERE / "references.bib")]

    ref = built / "reference.docx"
    make_reference_docx(pandoc, ref)
    docx = HERE / f"{OUT_STEM}.docx"
    subprocess.run([pandoc, str(md), *common, "--reference-doc", str(ref), "-o", str(docx)], check=True)
    shrink_docx_tables(docx)
    print("wrote", docx)

    css = built / "manuscript.css"
    css.write_text(CSS, encoding="utf-8")
    html = HERE / f"{OUT_STEM}.html"
    subprocess.run([pandoc, str(md), *common, "--standalone", "--mathml", "--embed-resources",
                    "--css", str(css), "--metadata", "lang=en", "-o", str(html)], check=True)
    print("wrote", html)

    if not args.no_pdf:
        browser = find_browser()
        if browser is None:
            raise SystemExit("no Chrome/Edge found for PDF printing; rerun with --no-pdf")
        pdf = HERE / f"{OUT_STEM}.pdf"
        prof = built / "chrome_profile"
        subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--user-data-dir={prof}", "--virtual-time-budget=15000",
                        f"--print-to-pdf={pdf}", html.as_uri()], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
        print("wrote", pdf)


if __name__ == "__main__":
    main()
